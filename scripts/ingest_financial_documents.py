"""Multi-format document ingestion for RAG system.

Supports: PDF, TXT, CSV, XLSX, DOCX, MD
Can be used as CLI script or imported for API routes.
"""
import asyncio
import sys
from pathlib import Path
import time
from typing import List, Dict, Any, Optional
from enum import Enum
import structlog

from src.rag_anything.client import RAGAnythingClient
from src.rag_anything.config import RAGConfig

# Configure structured logging
structlog.configure(
    processors=[
        structlog.processors.TimeStamper(fmt="iso"),
        structlog.processors.add_log_level,
        structlog.processors.JSONRenderer(),
    ],
    logger_factory=structlog.PrintLoggerFactory(),
)

logger = structlog.get_logger(__name__)


class FileFormat(str, Enum):
    """Supported file formats."""
    PDF = "pdf"
    TXT = "txt"
    CSV = "csv"
    XLSX = "xlsx"
    DOCX = "docx"
    MD = "md"


class DocumentProcessor:
    """Handles preprocessing of different file formats before RAG ingestion."""
    
    @staticmethod
    async def preprocess_txt(file_path: Path) -> str:
        """Read and return text file content.
        
        Args:
            file_path: Path to TXT file
            
        Returns:
            Text content
        """
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                content = f.read()
            return content
        except UnicodeDecodeError:
            # Fallback to latin-1 encoding
            with open(file_path, 'r', encoding='latin-1') as f:
                content = f.read()
            return content
    
    @staticmethod
    async def preprocess_csv(file_path: Path) -> str:
        """Convert CSV to formatted text.
        
        Args:
            file_path: Path to CSV file
            
        Returns:
            Formatted text representation
        """
        import pandas as pd
        
        try:
            df = pd.read_csv(file_path)
            
            # Create formatted text with metadata
            text_parts = [
                f"CSV Document: {file_path.name}",
                f"Rows: {len(df)}, Columns: {len(df.columns)}",
                f"\nColumn Names: {', '.join(df.columns.tolist())}",
                f"\nData:\n{df.to_string()}",
            ]
            
            # Add summary statistics for numeric columns
            numeric_cols = df.select_dtypes(include=['number']).columns
            if len(numeric_cols) > 0:
                text_parts.append(f"\n\nSummary Statistics:\n{df[numeric_cols].describe().to_string()}")
            
            return "\n".join(text_parts)
            
        except Exception as e:
            logger.error("csv_preprocessing_failed", file=file_path.name, error=str(e))
            raise
    
    @staticmethod
    async def preprocess_xlsx(file_path: Path) -> str:
        """Convert Excel to formatted text.
        
        Args:
            file_path: Path to XLSX file
            
        Returns:
            Formatted text representation
        """
        import pandas as pd
        
        try:
            # Read all sheets
            excel_file = pd.ExcelFile(file_path)
            sheet_names = excel_file.sheet_names
            
            text_parts = [
                f"Excel Document: {file_path.name}",
                f"Total Sheets: {len(sheet_names)}",
                f"Sheet Names: {', '.join(sheet_names)}\n"
            ]
            
            # Process each sheet
            for sheet_name in sheet_names:
                df = pd.read_excel(file_path, sheet_name=sheet_name)
                
                text_parts.append(f"\n{'='*70}")
                text_parts.append(f"SHEET: {sheet_name}")
                text_parts.append(f"{'='*70}")
                text_parts.append(f"Rows: {len(df)}, Columns: {len(df.columns)}")
                text_parts.append(f"Columns: {', '.join(df.columns.tolist())}")
                text_parts.append(f"\nData:\n{df.to_string()}")
                
                # Add summary for numeric columns
                numeric_cols = df.select_dtypes(include=['number']).columns
                if len(numeric_cols) > 0:
                    text_parts.append(f"\nSummary:\n{df[numeric_cols].describe().to_string()}")
            
            return "\n".join(text_parts)
            
        except Exception as e:
            logger.error("xlsx_preprocessing_failed", file=file_path.name, error=str(e))
            raise
    
    @staticmethod
    async def preprocess_docx(file_path: Path) -> str:
        """Extract text from DOCX.
        
        Args:
            file_path: Path to DOCX file
            
        Returns:
            Extracted text
        """
        try:
            import docx
            
            doc = docx.Document(file_path)
            
            text_parts = [f"Document: {file_path.name}\n"]
            
            # Extract paragraphs
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    text_parts.append(paragraph.text)
            
            # Extract tables
            if doc.tables:
                text_parts.append("\n\nTables:\n")
                for i, table in enumerate(doc.tables, 1):
                    text_parts.append(f"\nTable {i}:")
                    for row in table.rows:
                        row_text = " | ".join(cell.text.strip() for cell in row.cells)
                        text_parts.append(row_text)
            
            return "\n".join(text_parts)
            
        except Exception as e:
            logger.error("docx_preprocessing_failed", file=file_path.name, error=str(e))
            raise
    
    @staticmethod
    async def preprocess_markdown(file_path: Path) -> str:
        """Read markdown file.
        
        Args:
            file_path: Path to MD file
            
        Returns:
            Markdown content
        """
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                content = f.read()
            
            # Add metadata
            return f"Markdown Document: {file_path.name}\n\n{content}"
            
        except Exception as e:
            logger.error("markdown_preprocessing_failed", file=file_path.name, error=str(e))
            raise


class IngestionResult:
    """Result of document ingestion."""
    
    def __init__(self):
        self.total = 0
        self.successful = 0
        self.failed = 0
        self.skipped = 0
        self.total_time = 0.0
        self.results: List[Dict[str, Any]] = []
    
    def add_success(self, doc_name: str, elapsed: float):
        """Record successful ingestion."""
        self.successful += 1
        self.total_time += elapsed
        self.results.append({
            "document": doc_name,
            "status": "success",
            "elapsed": elapsed
        })
    
    def add_failure(self, doc_name: str, error: str, elapsed: float):
        """Record failed ingestion."""
        self.failed += 1
        self.total_time += elapsed
        self.results.append({
            "document": doc_name,
            "status": "failed",
            "error": error,
            "elapsed": elapsed
        })
    
    def add_skipped(self, doc_name: str, reason: str):
        """Record skipped document."""
        self.skipped += 1
        self.results.append({
            "document": doc_name,
            "status": "skipped",
            "reason": reason
        })
    
    def to_dict(self) -> Dict[str, Any]:
        """Convert to dictionary."""
        return {
            "total": self.total,
            "successful": self.successful,
            "failed": self.failed,
            "skipped": self.skipped,
            "total_time": round(self.total_time, 2),
            "avg_time": round(self.total_time / self.total, 2) if self.total > 0 else 0,
            "results": self.results
        }


async def ingest_documents(
    documents: List[Path],
    rag_client: Optional[RAGAnythingClient] = None,
    config: Optional[RAGConfig] = None,
    verbose: bool = True
) -> IngestionResult:
    """Ingest multiple documents into RAG system.
    
    This is the main function to call from API routes.
    
    Args:
        documents: List of document paths to ingest
        rag_client: Optional pre-initialized RAG client
        config: Optional RAG configuration
        verbose: Whether to print progress (False for API usage)
        
    Returns:
        IngestionResult with detailed results
    """
    result = IngestionResult()
    result.total = len(documents)
    
    if verbose:
        print("\n" + "="*70)
        print("  Document Ingestion Pipeline")
        print("="*70 + "\n")
    
    # Initialize RAG client if not provided
    if rag_client is None:
        if verbose:
            print("[1/3] Initializing RAG client...")
        
        config = config or RAGConfig()
        rag_client = RAGAnythingClient(config=config)
        
        try:
            await rag_client.initialize()
            if verbose:
                print("      ✓ RAG client initialized\n")
        except Exception as e:
            logger.error("initialization_failed", error=str(e))
            if verbose:
                print(f"      ✗ Initialization failed: {e}\n")
            return result
    
    # Process documents
    if verbose:
        print(f"[2/3] Processing {len(documents)} documents...")
        print("      " + "="*66 + "\n")
    
    processor = DocumentProcessor()
    
    for i, doc_path in enumerate(documents, 1):
        doc_name = doc_path.name
        file_format = doc_path.suffix.lower().lstrip('.')
        file_size_mb = doc_path.stat().st_size / (1024 * 1024)
        
        if verbose:
            print(f"      [{i}/{len(documents)}] {doc_name}")
            print(f"           Format: {file_format.upper()}, Size: {file_size_mb:.2f} MB", end="", flush=True)
        
        start_time = time.time()
        
        try:
            # Handle different file formats
            if file_format == FileFormat.PDF:
                # PDF: Direct processing with MinerU
                process_result = await rag_client.process_document(doc_path)
                
                if process_result.get("status") == "success":
                    elapsed = time.time() - start_time
                    result.add_success(doc_name, elapsed)
                    if verbose:
                        print(f" ... ✓ {elapsed:.1f}s")
                else:
                    error = process_result.get('error', 'Unknown error')
                    elapsed = time.time() - start_time
                    result.add_failure(doc_name, error, elapsed)
                    if verbose:
                        print(f" ... ✗ {error}")
            
            elif file_format == FileFormat.TXT:
                # TXT: Read and insert as text
                content = await processor.preprocess_txt(doc_path)
                
                # Insert text directly into RAG
                if hasattr(rag_client._rag, 'lightrag') and rag_client._rag.lightrag:
                    await rag_client._rag.lightrag.ainsert(content)
                    elapsed = time.time() - start_time
                    result.add_success(doc_name, elapsed)
                    if verbose:
                        print(f" ... ✓ {elapsed:.1f}s")
                else:
                    raise RuntimeError("LightRAG not initialized")
            
            elif file_format == FileFormat.CSV:
                # CSV: Convert to text and insert
                content = await processor.preprocess_csv(doc_path)
                
                if hasattr(rag_client._rag, 'lightrag') and rag_client._rag.lightrag:
                    await rag_client._rag.lightrag.ainsert(content)
                    elapsed = time.time() - start_time
                    result.add_success(doc_name, elapsed)
                    if verbose:
                        print(f" ... ✓ {elapsed:.1f}s")
                else:
                    raise RuntimeError("LightRAG not initialized")
            
            elif file_format == FileFormat.XLSX:
                # XLSX: Convert to text and insert
                content = await processor.preprocess_xlsx(doc_path)
                
                if hasattr(rag_client._rag, 'lightrag') and rag_client._rag.lightrag:
                    await rag_client._rag.lightrag.ainsert(content)
                    elapsed = time.time() - start_time
                    result.add_success(doc_name, elapsed)
                    if verbose:
                        print(f" ... ✓ {elapsed:.1f}s")
                else:
                    raise RuntimeError("LightRAG not initialized")
            
            elif file_format == FileFormat.DOCX:
                # DOCX: Extract text and insert
                content = await processor.preprocess_docx(doc_path)
                
                if hasattr(rag_client._rag, 'lightrag') and rag_client._rag.lightrag:
                    await rag_client._rag.lightrag.ainsert(content)
                    elapsed = time.time() - start_time
                    result.add_success(doc_name, elapsed)
                    if verbose:
                        print(f" ... ✓ {elapsed:.1f}s")
                else:
                    raise RuntimeError("LightRAG not initialized")
            
            elif file_format == FileFormat.MD:
                # Markdown: Read and insert
                content = await processor.preprocess_markdown(doc_path)
                
                if hasattr(rag_client._rag, 'lightrag') and rag_client._rag.lightrag:
                    await rag_client._rag.lightrag.ainsert(content)
                    elapsed = time.time() - start_time
                    result.add_success(doc_name, elapsed)
                    if verbose:
                        print(f" ... ✓ {elapsed:.1f}s")
                else:
                    raise RuntimeError("LightRAG not initialized")
            
            else:
                # Unsupported format
                result.add_skipped(doc_name, f"Unsupported format: {file_format}")
                if verbose:
                    print(f" ... ⊘ Unsupported format")
        
        except Exception as e:
            elapsed = time.time() - start_time
            error_msg = str(e)
            result.add_failure(doc_name, error_msg, elapsed)
            if verbose:
                print(f" ... ✗ {error_msg}")
            logger.error("document_processing_failed", document=doc_name, error=error_msg)
        
        # Progress update
        if verbose:
            remaining = len(documents) - i
            avg_time = result.total_time / i
            est_remaining = remaining * avg_time
            
            print(f"           Progress: [{i}/{len(documents)}] | " +
                  f"Success: {result.successful} | Failed: {result.failed} | " +
                  f"Skipped: {result.skipped} | " +
                  f"ETA: {est_remaining:.0f}s")
            print()
    
    # Summary
    if verbose:
        print("[3/3] Ingestion Complete!\n")
        print("="*70)
        print("  SUMMARY")
        print("="*70)
        print(f"  Total documents: {result.total}")
        print(f"  Successful: {result.successful}")
        print(f"  Failed: {result.failed}")
        print(f"  Skipped: {result.skipped}")
        print(f"  Total time: {result.total_time:.1f}s")
        if result.total > 0:
            print(f"  Average per doc: {result.total_time/result.total:.1f}s")
        print("="*70 + "\n")
    
    return result


async def ingest_from_directory(
    directory: Path,
    file_patterns: Optional[List[str]] = None,
    verbose: bool = True
) -> IngestionResult:
    """Ingest all documents from a directory.
    
    Args:
        directory: Directory containing documents
        file_patterns: Optional list of glob patterns (default: all supported formats)
        verbose: Whether to print progress
        
    Returns:
        IngestionResult with detailed results
    """
    if not directory.exists():
        raise FileNotFoundError(f"Directory not found: {directory}")
    
    # Default patterns for all supported formats
    if file_patterns is None:
        file_patterns = ["*.pdf", "*.txt", "*.csv", "*.xlsx", "*.docx", "*.md"]
    
    # Collect all matching files
    documents = []
    for pattern in file_patterns:
        documents.extend(directory.glob(pattern))
    
    if not documents:
        logger.warning("no_documents_found", directory=str(directory))
        if verbose:
            print(f"No documents found in {directory}")
        return IngestionResult()
    
    return await ingest_documents(documents, verbose=verbose)


# CLI Entry Point
async def main():
    """CLI entry point for document ingestion."""
    print("\nStarting document ingestion pipeline...")
    
    # Use directory from command line or default
    if len(sys.argv) > 1:
        doc_dir = Path(sys.argv[1])
    else:
        doc_dir = Path("data/test1")
    
    result = await ingest_from_directory(doc_dir, verbose=True)
    
    # Test retrieval if successful
    if result.successful > 0:
        print("\n" + "="*70)
        print("  TESTING RETRIEVAL")
        print("="*70 + "\n")
        
        config = RAGConfig()
        rag_client = RAGAnythingClient(config=config)
        await rag_client.initialize()
        
        test_query = "summarize the key financial metrics"
        print(f"Test query: '{test_query}'...")
        
        try:
            query_result = await rag_client.query(
                query_text=test_query,
                mode="hybrid"
            )
            
            context = query_result.get("context", "")
            
            if context and len(context) > 100:
                print(f"\n✓ Retrieved {len(context)} characters")
                print("\nSample context:")
                print("-" * 70)
                preview = context[:500] + "..." if len(context) > 500 else context
                print(preview)
                print("-" * 70)
                print("\n✓ RAG retrieval working!\n")
            else:
                print("\n⚠ Context empty or too short\n")
                
        except Exception as e:
            print(f"\n✗ Test query failed: {e}\n")


if __name__ == "__main__":
    asyncio.run(main())