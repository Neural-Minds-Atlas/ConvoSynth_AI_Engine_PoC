# File: rag_anything/pipeline/document_processor.py
# ============================================================================
"""Multi-format document processing with MinerU integration."""

from pathlib import Path
from typing import Dict, Any, Optional
from enum import Enum
import structlog

logger = structlog.get_logger(__name__)


class FileFormat(str, Enum):
    """Supported file formats."""
    PDF = "pdf"
    TXT = "txt"
    CSV = "csv"
    XLSX = "xlsx"
    DOCX = "docx"
    MD = "md"


class DocumentProcessor:
    """Handles preprocessing of different file formats."""
    
    def __init__(self):
        """Initialize document processor."""
        self.logger = logger.bind(component="document_processor")
    
    async def process(self, file_path: Path) -> Dict[str, Any]:
        """Process document based on format.
        
        Args:
            file_path: Path to document
            
        Returns:
            Processing result with content and metadata
        """
        file_format = file_path.suffix.lower().lstrip('.')
        
        try:
            if file_format == FileFormat.PDF:
                return await self._process_pdf(file_path)
            elif file_format == FileFormat.TXT:
                return await self._process_txt(file_path)
            elif file_format == FileFormat.CSV:
                return await self._process_csv(file_path)
            elif file_format == FileFormat.XLSX:
                return await self._process_xlsx(file_path)
            elif file_format == FileFormat.DOCX:
                return await self._process_docx(file_path)
            elif file_format == FileFormat.MD:
                return await self._process_markdown(file_path)
            else:
                raise ValueError(f"Unsupported format: {file_format}")
                
        except Exception as e:
            self.logger.error("processing_failed", file=file_path.name, error=str(e))
            raise
    
    async def _process_pdf(self, file_path: Path) -> Dict[str, Any]:
        """Process PDF with MinerU (handled by RAG-Anything)."""
        return {
            "format": "pdf",
            "path": str(file_path),
            "use_mineru": True  # Signal to use MinerU parser
        }
    
    async def _process_txt(self, file_path: Path) -> Dict[str, Any]:
        """Process plain text file."""
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                content = f.read()
        except UnicodeDecodeError:
            with open(file_path, 'r', encoding='latin-1') as f:
                content = f.read()
        
        return {
            "format": "txt",
            "content": content,
            "metadata": {
                "filename": file_path.name,
                "size": file_path.stat().st_size
            }
        }
    
    async def _process_csv(self, file_path: Path) -> Dict[str, Any]:
        """Process CSV file."""
        import pandas as pd
        
        df = pd.read_csv(file_path)
        
        # Create formatted text representation
        text_parts = [
            f"CSV Document: {file_path.name}",
            f"Rows: {len(df)}, Columns: {len(df.columns)}",
            f"\nColumn Names: {', '.join(df.columns.tolist())}",
            f"\nData:\n{df.to_string()}",
        ]
        
        # Add summary statistics
        numeric_cols = df.select_dtypes(include=['number']).columns
        if len(numeric_cols) > 0:
            text_parts.append(f"\n\nSummary Statistics:\n{df[numeric_cols].describe().to_string()}")
        
        return {
            "format": "csv",
            "content": "\n".join(text_parts),
            "metadata": {
                "filename": file_path.name,
                "rows": len(df),
                "columns": len(df.columns),
                "column_names": df.columns.tolist()
            }
        }
    
    async def _process_xlsx(self, file_path: Path) -> Dict[str, Any]:
        """Process Excel file."""
        import pandas as pd
        
        excel_file = pd.ExcelFile(file_path)
        sheet_names = excel_file.sheet_names
        
        text_parts = [
            f"Excel Document: {file_path.name}",
            f"Total Sheets: {len(sheet_names)}",
            f"Sheet Names: {', '.join(sheet_names)}\n"
        ]
        
        for sheet_name in sheet_names:
            df = pd.read_excel(file_path, sheet_name=sheet_name)
            
            text_parts.append(f"\n{'='*70}")
            text_parts.append(f"SHEET: {sheet_name}")
            text_parts.append(f"{'='*70}")
            text_parts.append(f"Rows: {len(df)}, Columns: {len(df.columns)}")
            text_parts.append(f"Columns: {', '.join(df.columns.tolist())}")
            text_parts.append(f"\nData:\n{df.to_string()}")
            
            numeric_cols = df.select_dtypes(include=['number']).columns
            if len(numeric_cols) > 0:
                text_parts.append(f"\nSummary:\n{df[numeric_cols].describe().to_string()}")
        
        return {
            "format": "xlsx",
            "content": "\n".join(text_parts),
            "metadata": {
                "filename": file_path.name,
                "sheets": sheet_names,
                "total_sheets": len(sheet_names)
            }
        }
    
    async def _process_docx(self, file_path: Path) -> Dict[str, Any]:
        """Process Word document."""
        import docx
        
        doc = docx.Document(file_path)
        text_parts = [f"Document: {file_path.name}\n"]
        
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text_parts.append(paragraph.text)
        
        if doc.tables:
            text_parts.append("\n\nTables:\n")
            for i, table in enumerate(doc.tables, 1):
                text_parts.append(f"\nTable {i}:")
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells)
                    text_parts.append(row_text)
        
        return {
            "format": "docx",
            "content": "\n".join(text_parts),
            "metadata": {
                "filename": file_path.name,
                "paragraphs": len(doc.paragraphs),
                "tables": len(doc.tables)
            }
        }
    
    async def _process_markdown(self, file_path: Path) -> Dict[str, Any]:
        """Process Markdown file."""
        with open(file_path, 'r', encoding='utf-8') as f:
            content = f.read()
        
        return {
            "format": "md",
            "content": f"Markdown Document: {file_path.name}\n\n{content}",
            "metadata": {
                "filename": file_path.name
            }
        }